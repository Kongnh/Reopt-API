import unittest
from unittest import TestCase

from openpyxl import Workbook

from proforma_vietnam.tools.compare_workbooks import compare_workbooks


def _write(tmp_path, name, rows, sheet="Sheet1"):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    path = tmp_path / name
    workbook.save(path)
    return path


class CompareWorkbooksTests(TestCase):

    def setUp(self):
        import tempfile
        from pathlib import Path
        self._dir = tempfile.TemporaryDirectory()
        self.tmp = Path(self._dir.name)
        self.addCleanup(self._dir.cleanup)

    def test_identical_workbooks_report_no_differences(self):
        a = _write(self.tmp, "a.xlsx", [["x", 1], ["y", 2]])
        b = _write(self.tmp, "b.xlsx", [["x", 1], ["y", 2]])

        self.assertEqual(compare_workbooks(a, b), [])

    def test_differing_cell_is_reported_with_sheet_and_coordinates(self):
        a = _write(self.tmp, "a.xlsx", [["x", 1]])
        b = _write(self.tmp, "b.xlsx", [["x", 2]])

        differences = compare_workbooks(a, b)

        self.assertEqual(len(differences), 1)
        self.assertIn("Sheet1", differences[0])
        self.assertIn("row 1", differences[0])
        self.assertIn("col 2", differences[0])

    def test_ignored_substring_suppresses_the_difference(self):
        # DEFAULT_IGNORE_SUBSTRINGS is now empty (see NarrowedIgnoreRulesTests),
        # so this exercises the ignore_substrings parameter explicitly rather
        # than relying on a default that no longer ignores anything.
        a = _write(self.tmp, "a.xlsx", [["prepared 2026-07-07"]])
        b = _write(self.tmp, "b.xlsx", [["prepared 2026-09-04"]])

        self.assertEqual(
            compare_workbooks(a, b, ignore_substrings=("prepared ",)), [])

    def test_ignored_substring_must_match_both_sides(self):
        a = _write(self.tmp, "a.xlsx", [["prepared 2026-07-07"]])
        b = _write(self.tmp, "b.xlsx", [["something else"]])

        self.assertEqual(
            len(compare_workbooks(a, b, ignore_substrings=("prepared ",))), 1)

    def test_sheet_name_mismatch_is_reported(self):
        a = _write(self.tmp, "a.xlsx", [["x"]], sheet="Alpha")
        b = _write(self.tmp, "b.xlsx", [["x"]], sheet="Beta")

        differences = compare_workbooks(a, b)

        self.assertEqual(len(differences), 1)
        self.assertIn("sheet names differ", differences[0])



class VolatileRowTests(TestCase):
    """The run-date row must not fail the gate once the clock rolls over."""

    def setUp(self):
        import tempfile
        from pathlib import Path
        self._dir = tempfile.TemporaryDirectory()
        self.tmp = Path(self._dir.name)
        self.addCleanup(self._dir.cleanup)

    def test_a_differing_run_date_is_not_a_difference(self):
        a = _write(self.tmp, "a.xlsx",
                   [["Report prepared", "2026-09-04"], ["Analysis period", 25]],
                   sheet="Assumptions")
        b = _write(self.tmp, "b.xlsx",
                   [["Report prepared", "2026-09-05"], ["Analysis period", 25]],
                   sheet="Assumptions")

        self.assertEqual(compare_workbooks(a, b), [])

    def test_a_real_change_on_another_row_still_reports(self):
        a = _write(self.tmp, "a.xlsx",
                   [["Report prepared", "2026-09-04"], ["Analysis period", 25]],
                   sheet="Assumptions")
        b = _write(self.tmp, "b.xlsx",
                   [["Report prepared", "2026-09-05"], ["Analysis period", 20]],
                   sheet="Assumptions")

        self.assertEqual(len(compare_workbooks(a, b)), 1)


class NarrowedIgnoreRulesTests(unittest.TestCase):

    def _workbook(self, path, rows):
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=column_index, value=value)
        workbook.save(path)
        return path

    def test_a_changed_subtitle_is_now_visible(self):
        import tempfile, os
        directory = tempfile.mkdtemp()
        a = self._workbook(os.path.join(directory, "a.xlsx"),
                           [["Case prepared for Factory A"]])
        b = self._workbook(os.path.join(directory, "b.xlsx"),
                           [["Case prepared for Factory B"]])
        self.assertEqual(len(compare_workbooks(a, b)), 1)

    def test_the_report_prepared_date_row_is_still_ignored(self):
        import tempfile, os
        directory = tempfile.mkdtemp()
        a = self._workbook(os.path.join(directory, "a.xlsx"),
                           [["Report prepared", "2026-09-04"]])
        b = self._workbook(os.path.join(directory, "b.xlsx"),
                           [["Report prepared", "2026-09-05"]])
        self.assertEqual(compare_workbooks(a, b), [])


class PreparedDateNormalizationTests(TestCase):
    """The "prepared <date>" token interpolated into the Cover and Executive
    Summary subtitles is neutralised on both sides before comparison, but
    every other character of the subtitle is still compared for real."""

    def setUp(self):
        import tempfile
        from pathlib import Path
        self._dir = tempfile.TemporaryDirectory()
        self.tmp = Path(self._dir.name)
        self.addCleanup(self._dir.cleanup)

    def test_subtitles_differing_only_by_prepared_date_compare_equal(self):
        a = _write(self.tmp, "a.xlsx",
                   [["Factory A  ·  prepared 2026-09-04  ·  REopt dispatch"]])
        b = _write(self.tmp, "b.xlsx",
                   [["Factory A  ·  prepared 2026-09-05  ·  REopt dispatch"]])

        self.assertEqual(compare_workbooks(a, b), [])

    def test_same_prepared_date_but_different_case_name_still_reports(self):
        # This is the regression Task 11 exists to catch: a substring/whole-
        # row ignore rule would hide this too, which is how a Vietnam-titled
        # workbook reached a Thailand client.
        a = _write(self.tmp, "a.xlsx",
                   [["Factory A  ·  prepared 2026-09-04  ·  REopt dispatch"]])
        b = _write(self.tmp, "b.xlsx",
                   [["Factory B  ·  prepared 2026-09-04  ·  REopt dispatch"]])

        self.assertEqual(len(compare_workbooks(a, b)), 1)

    def test_prepared_date_and_case_name_both_differing_still_reports(self):
        a = _write(self.tmp, "a.xlsx",
                   [["Factory A  ·  prepared 2026-09-04  ·  REopt dispatch"]])
        b = _write(self.tmp, "b.xlsx",
                   [["Factory B  ·  prepared 2026-09-05  ·  REopt dispatch"]])

        self.assertEqual(len(compare_workbooks(a, b)), 1)

    def test_bare_date_not_preceded_by_prepared_is_compared_normally(self):
        # A genuine date change with no "prepared " prefix must not be
        # silently swallowed by the normalization.
        a = _write(self.tmp, "a.xlsx", [["2026-09-04"]])
        b = _write(self.tmp, "b.xlsx", [["2026-09-05"]])

        self.assertEqual(len(compare_workbooks(a, b)), 1)
