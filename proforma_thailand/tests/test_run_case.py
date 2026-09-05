from unittest import TestCase

from proforma_thailand.run_case import summarize_results


class SummarizeResultsTests(TestCase):

    def _results(self, pv_kw=1000.0, pv_kwh=1_500_000.0, grid_kwh=5_000_000.0):
        return {
            "outputs": {
                "PV": {"size_kw": pv_kw, "annual_energy_produced_kwh": pv_kwh},
                "ElectricStorage": {"size_kw": 0.0, "size_kwh": 0.0},
                "ElectricLoad": {"annual_calculated_kwh": 6_500_000.0},
                "ElectricUtility": {"annual_energy_supplied_kwh": grid_kwh},
            }
        }

    def test_grid_offset_is_one_minus_grid_over_load(self):
        summary = summarize_results(self._results(), extras={})

        self.assertAlmostEqual(
            summary["grid_offset_fraction"], 1 - 5_000_000.0 / 6_500_000.0, places=6
        )

    def test_sizes_are_carried_through(self):
        summary = summarize_results(self._results(), extras={})

        self.assertEqual(summary["pv_kw"], 1000.0)
        self.assertEqual(summary["bess_kw"], 0.0)

    def test_zero_load_does_not_divide_by_zero(self):
        results = self._results()
        results["outputs"]["ElectricLoad"]["annual_calculated_kwh"] = 0.0

        summary = summarize_results(results, extras={})

        self.assertEqual(summary["grid_offset_fraction"], 0.0)

    def test_report_extras_are_merged_into_the_summary(self):
        summary = summarize_results(
            self._results(),
            extras={"power_factor_compensation_kvar": 168.24,
                    "power_factor_mitigation_cost_usd": 6000.0},
        )

        self.assertAlmostEqual(
            summary["power_factor_compensation_kvar"], 168.24, places=3
        )
        self.assertEqual(summary["power_factor_mitigation_cost_usd"], 6000.0)

    def test_missing_extras_default_to_zero(self):
        summary = summarize_results(self._results(), extras={})

        self.assertEqual(summary["power_factor_compensation_kvar"], 0.0)


class PollCompletenessTests(TestCase):
    """status can read optimal while outputs are still being written."""

    def test_an_optimal_run_without_outputs_is_not_complete(self):
        from proforma_thailand.run_case import _is_complete

        self.assertFalse(_is_complete({
            "status": "optimal",
            "outputs": {"Financial": {}, "ElectricTariff": {}},
        }))

    def test_an_optimal_run_with_outputs_is_complete(self):
        from proforma_thailand.run_case import _is_complete

        self.assertTrue(_is_complete({
            "status": "optimal",
            "outputs": {"ElectricLoad": {"annual_calculated_kwh": 1.0}},
        }))

    def test_an_error_run_is_complete_without_outputs(self):
        from proforma_thailand.run_case import _is_complete

        self.assertTrue(_is_complete({"status": "error", "outputs": {}}))


class PlaceholderGuardTests(TestCase):

    def _workbook(self, cells):
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for index, value in enumerate(cells, start=1):
            sheet.cell(row=index, column=1, value=value)
        return workbook

    def test_marked_workbook_passes(self):
        from proforma_thailand.run_case import assert_placeholders_disclosed

        workbook = self._workbook([
            "Equity IRR", "PLACEHOLDER - pending Keen confirmation",
        ])
        assert_placeholders_disclosed(workbook)

    def test_unmarked_workbook_raises(self):
        from proforma_thailand.run_case import assert_placeholders_disclosed

        workbook = self._workbook(["Equity IRR", "all inputs confirmed"])
        with self.assertRaises(RuntimeError) as caught:
            assert_placeholders_disclosed(workbook)
        self.assertIn("placeholder", str(caught.exception).lower())

    def test_workbook_with_no_headline_metric_passes(self):
        from proforma_thailand.run_case import assert_placeholders_disclosed

        workbook = self._workbook(["Dispatch", "Interval"])
        assert_placeholders_disclosed(workbook)
