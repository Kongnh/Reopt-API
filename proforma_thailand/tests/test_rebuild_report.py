"""Cover the offline Thailand workbook rebuild."""
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from proforma_thailand.rebuild_report import rebuild_report

CASE_DIR = Path("outputs/thailand_case/rofu_thailand/case_1")


class RebuildReportTests(unittest.TestCase):
    def _scratch_case(self):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        for name in ("results.json", "assumptions.json", "case.json"):
            source = CASE_DIR / name
            if source.exists():
                shutil.copy2(source, tmp / name)
        return tmp

    def test_rebuild_writes_a_workbook_named_for_the_run(self):
        case_dir = self._scratch_case()
        results = json.loads((case_dir / "results.json").read_text(encoding="utf-8"))

        out_path = rebuild_report(case_dir)

        self.assertTrue(out_path.exists())
        self.assertEqual(
            out_path.name, "thailand_report_{}.xlsx".format(results["run_uuid"])
        )

    def test_prepared_on_pins_the_date_so_two_rebuilds_match(self):
        """Without a pinned date the workbook carries date.today() in three
        cells, which would make every baseline comparison fail a day later."""
        from openpyxl import load_workbook

        first = load_workbook(rebuild_report(self._scratch_case(), prepared_on="2026-01-01"))
        second = load_workbook(rebuild_report(self._scratch_case(), prepared_on="2026-01-01"))

        sheet_one = first["Executive Summary"]
        sheet_two = second["Executive Summary"]
        self.assertEqual(
            [row for row in sheet_one.iter_rows(values_only=True)],
            [row for row in sheet_two.iter_rows(values_only=True)],
        )
        joined = " ".join(
            str(value)
            for row in sheet_one.iter_rows(values_only=True)
            for value in row
            if isinstance(value, str)
        )
        self.assertIn("prepared 2026-01-01", joined)


if __name__ == "__main__":
    unittest.main()
