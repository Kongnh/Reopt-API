"""Turn the Rofu 15-minute meter workbook into a REopt load array.

The source sheet ("Raw Total Consumption") covers 416 unbroken days from
2025-06-01 to 2026-07-21 with 96 intervals every day and no nulls. It also
carries a DayType column that reproduces PEA's own billing day classification
(Saturdays, Sundays and Thai public holidays are all "Holiday"), verified by
reconciling calendar June 2025 against that month's invoice to within 0.003%.

The analysis year is a synthetic calendar year - January to June taken from
2026, July to December from 2025 - so that load, tariff and the calendar-year
PVWatts production factor all index identically from 1 January. A contiguous
July-to-June window would pair wet-season load with dry-season irradiance.
"""

from datetime import date, timedelta

INTERVALS_PER_DAY = 96
MINUTES_PER_INTERVAL = 15
SHEET_NAME = "Raw Total Consumption"


def extract_intervals(xlsm_path):
    """Read the meter sheet into interval dicts."""
    from openpyxl import load_workbook

    workbook = load_workbook(xlsm_path, read_only=True, data_only=True)
    worksheet = workbook[SHEET_NAME]
    rows = []
    for index, row in enumerate(worksheet.iter_rows(values_only=True)):
        if index == 0 or row[0] is None:
            continue
        timestamp = row[0]
        clock = row[2]
        if isinstance(clock, timedelta):
            # openpyxl cannot express a time-of-day of exactly 24:00 as a
            # datetime.time, so the interval labelled 00:00 (end of day,
            # Excel serial 1.0) comes back as timedelta(days=1) instead.
            end_minute = int(clock.total_seconds() // 60)
        else:
            end_minute = clock.hour * 60 + clock.minute
            if end_minute == 0:
                end_minute = 24 * 60
        rows.append(
            {
                "date": timestamp.date(),
                "end_minute": end_minute,
                "day_type": row[4],
                "kwh": float(row[9] or 0.0),
            }
        )
    return rows


def build_calendar_year(intervals, calendar_year_months):
    """Return ``(loads_kw, all_off_peak_dates, qa)`` for the synthetic year."""
    by_date = {}
    for interval in intervals:
        by_date.setdefault(interval["date"], []).append(interval)

    wanted = []
    for year, month in calendar_year_months:
        wanted.extend(_days_in_month(year, month))

    missing = [day for day in wanted if day not in by_date]
    if missing:
        raise ValueError(
            "Load data is missing {} day(s), first {}.".format(
                len(missing), missing[0].isoformat()
            )
        )

    loads_kw = []
    all_off_peak_dates = set()
    total_kwh = 0.0
    for day in wanted:
        day_rows = sorted(by_date[day], key=lambda row: row["end_minute"])
        if len(day_rows) != INTERVALS_PER_DAY:
            raise ValueError(
                "{} has {} intervals, expected {}.".format(
                    day.isoformat(), len(day_rows), INTERVALS_PER_DAY
                )
            )
        if any(row["day_type"] == "Holiday" for row in day_rows):
            all_off_peak_dates.add(day)
        for row in day_rows:
            total_kwh += row["kwh"]
            loads_kw.append(row["kwh"] * (60.0 / MINUTES_PER_INTERVAL))

    qa = {
        "days": len(wanted),
        "missing_days": [],
        "total_kwh": total_kwh,
        "peak_kw": max(loads_kw),
    }
    return loads_kw, all_off_peak_dates, qa


def _days_in_month(year, month):
    current = date(year, month, 1)
    days = []
    while current.month == month:
        days.append(current)
        current += timedelta(days=1)
    return days
